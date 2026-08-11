"""Local-only document extraction and deterministic BioJob exports."""

from __future__ import annotations

import re
import shutil
from collections.abc import Iterable, Mapping, Sequence
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader
import xlsxwriter


MAX_PROFILE_DOCUMENT_BYTES = 10 * 1024 * 1024
MAX_IMPORTED_FACTS = 40
WORKBOOK_SHEETS = ("投递总表", "候选岗位", "本周行动", "数据字典")


def extract_profile_lines(path: Path) -> list[str]:
    suffix = path.suffix.casefold()
    if suffix == ".docx":
        document = Document(str(path))
        raw = [paragraph.text for paragraph in document.paragraphs]
        raw.extend(
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        )
    elif suffix == ".pdf":
        reader = PdfReader(path)
        raw = [page.extract_text() or "" for page in reader.pages]
    else:
        raise ValueError("profile document must be a .docx or .pdf file")

    lines: list[str] = []
    for block in raw:
        for candidate in block.splitlines():
            line = re.sub(r"\s+", " ", candidate).strip(" \t•·-—")
            if len(line) >= 2 and line not in lines:
                lines.append(line[:1000])
            if len(lines) >= MAX_IMPORTED_FACTS:
                return lines
    return lines


def copy_profile_document(source: Path, destination_dir: Path, digest: str) -> Path:
    destination_dir.mkdir(parents=True, exist_ok=True)
    safe_stem = _safe_filename(source.stem, fallback="resume")
    destination = (
        destination_dir / f"{safe_stem}-{digest[:12]}{source.suffix.casefold()}"
    )
    if not destination.exists():
        shutil.copy2(source, destination)
    return destination


def write_resume_docx(
    output: Path,
    *,
    job: Mapping[str, Any],
    facts: Sequence[Mapping[str, Any]],
) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    styles = document.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"].paragraph_format.space_after = Pt(2)

    name = _first_fact_value(facts, {"name", "full_name", "姓名"}) or "个人简历"
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(name)
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(28, 45, 39)

    company = _company_name(job)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(f"目标岗位｜{company} · {job.get('title', '')}")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(70, 96, 86)

    by_category: dict[str, list[Mapping[str, Any]]] = {}
    for fact in facts:
        by_category.setdefault(str(fact.get("category") or "其他经历"), []).append(fact)

    labels = {
        "basic": "基本信息",
        "education": "教育背景",
        "skills": "专业技能",
        "experience": "实践经历",
        "experiments": "实验经历",
        "projects": "项目经历",
        "certificates": "证书与工具",
    }
    preferred = [
        "basic",
        "education",
        "skills",
        "experiments",
        "experience",
        "projects",
        "certificates",
    ]
    ordered = [key for key in preferred if key in by_category]
    ordered.extend(key for key in by_category if key not in ordered)
    for category in ordered:
        _add_section_heading(document, labels.get(category, category))
        for fact in by_category[category]:
            value = _display_value(fact.get("value"))
            if not value:
                continue
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.left_indent = Inches(0.18)
            paragraph.paragraph_format.first_line_indent = Inches(-0.12)
            paragraph.add_run(value)

    _add_section_heading(document, "岗位定位")
    paragraph = document.add_paragraph()
    paragraph.add_run("应聘方向：").bold = True
    paragraph.add_run(str(job.get("direction") or job.get("title") or ""))
    if job.get("city"):
        paragraph.add_run(f"　工作地点：{job['city']}")

    core = document.core_properties
    core.title = f"{company}-{job.get('title', '')}-定制简历"
    core.subject = "BioJob Agent evidence-locked resume"
    core.author = "BioJob Agent"
    document.save(str(output))


def write_application_workbook(
    output: Path,
    *,
    jobs: Sequence[Mapping[str, Any]],
    candidates: Sequence[Mapping[str, Any]],
    resume_by_job: Mapping[str, str],
    match_report_by_job: Mapping[str, str],
) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = xlsxwriter.Workbook(output)
    workbook.set_properties({"title": "BioJob 求职投递表", "author": "BioJob Agent"})
    palette = {
        "ink": "#17342B",
        "green": "#2F6B57",
        "mint": "#E8F2ED",
        "sand": "#F6F1E8",
        "line": "#D7E0DB",
        "white": "#FFFFFF",
        "amber": "#FFF0C2",
    }
    title = workbook.add_format({
        "bold": True,
        "font_size": 18,
        "font_color": palette["ink"],
    })
    subtitle = workbook.add_format({"font_color": "#60736B", "font_size": 9})
    header = workbook.add_format({
        "bold": True,
        "font_color": palette["white"],
        "bg_color": palette["green"],
        "align": "center",
        "valign": "vcenter",
    })
    text = workbook.add_format({"font_color": palette["ink"], "valign": "top"})
    wrap = workbook.add_format({
        "font_color": palette["ink"],
        "valign": "top",
        "text_wrap": True,
    })
    link = workbook.add_format({"font_color": "#176B87", "underline": True})
    note = workbook.add_format({
        "bg_color": palette["sand"],
        "font_color": "#5B5143",
        "text_wrap": True,
    })
    status_formats = {
        "preparing": workbook.add_format({
            "bg_color": palette["amber"],
            "font_color": palette["ink"],
        }),
        "applied": workbook.add_format({
            "bg_color": palette["mint"],
            "font_color": palette["ink"],
        }),
        "offer": workbook.add_format({"bg_color": "#D9EFD9", "font_color": "#24552F"}),
        "rejected": workbook.add_format({
            "bg_color": "#FBE4E1",
            "font_color": "#8A3028",
        }),
    }

    total_headers = [
        "公司",
        "岗位",
        "城市",
        "方向",
        "投递状态",
        "匹配分",
        "推荐结论",
        "发布时间",
        "截止日期",
        "下次跟进",
        "JD",
        "投递入口",
        "招聘官网",
        "匹配报告",
        "定制简历",
        "备注",
    ]
    total = _sheet_frame(workbook, "投递总表", total_headers, title, subtitle, header)
    for row_index, job in enumerate(jobs, start=3):
        app = job.get("application") or {}
        status = str(app.get("status") or "")
        match = job.get("latest_match") or {}
        values = [
            _company_name(job),
            job.get("title"),
            job.get("city"),
            job.get("direction"),
            status,
            match.get("score"),
            match.get("recommendation"),
            job.get("published_at"),
            job.get("deadline_at"),
            app.get("next_follow_up_at"),
            None,
            None,
            None,
            None,
            None,
            app.get("notes") or job.get("notes"),
        ]
        for column, value in enumerate(values):
            cell_format = (
                status_formats.get(status, text)
                if column == 4
                else (wrap if column == 15 else text)
            )
            total.write(row_index, column, value, cell_format)
        _write_http_link(total, row_index, 10, job.get("detail_url"), "打开JD", link)
        _write_http_link(total, row_index, 11, job.get("apply_url"), "去投递", link)
        _write_http_link(total, row_index, 12, job.get("careers_url"), "招聘官网", link)
        _write_file_link(
            total,
            row_index,
            13,
            match_report_by_job.get(str(job.get("id"))),
            "查看报告",
            link,
        )
        _write_file_link(
            total,
            row_index,
            14,
            resume_by_job.get(str(job.get("id"))),
            "打开简历",
            link,
        )
    _finish_table(
        total,
        len(jobs),
        len(total_headers),
        [18, 22, 11, 14, 12, 10, 16, 12, 12, 18, 10, 10, 11, 11, 11, 28],
    )

    candidate_headers = [
        "公司",
        "岗位",
        "城市",
        "方向",
        "决策",
        "匹配分",
        "推荐结论",
        "JD",
        "投递入口",
        "招聘官网",
    ]
    candidate_sheet = _sheet_frame(
        workbook, "候选岗位", candidate_headers, title, subtitle, header
    )
    for row_index, candidate in enumerate(candidates, start=3):
        match = candidate.get("match") or {}
        values = [
            _company_name(candidate),
            candidate.get("title"),
            candidate.get("city"),
            candidate.get("direction"),
            candidate.get("decision"),
            match.get("score"),
            match.get("recommendation"),
            None,
            None,
            None,
        ]
        candidate_sheet.write_row(row_index, 0, values, text)
        _write_http_link(
            candidate_sheet, row_index, 7, candidate.get("detail_url"), "打开JD", link
        )
        _write_http_link(
            candidate_sheet, row_index, 8, candidate.get("apply_url"), "去投递", link
        )
        _write_http_link(
            candidate_sheet,
            row_index,
            9,
            candidate.get("careers_url"),
            "招聘官网",
            link,
        )
    _finish_table(
        candidate_sheet,
        len(candidates),
        len(candidate_headers),
        [18, 22, 11, 14, 10, 10, 16, 10, 10, 11],
    )

    action_headers = ["优先级", "行动", "公司", "岗位", "截止日期", "跟进时间", "入口"]
    action_sheet = _sheet_frame(
        workbook, "本周行动", action_headers, title, subtitle, header
    )
    actions = [
        job
        for job in jobs
        if (job.get("application") or {}).get("status")
        in {"preparing", "applied", "assessment", "interview"}
    ]
    for row_index, job in enumerate(actions, start=3):
        app = job.get("application") or {}
        status = app.get("status")
        action = "完善材料并投递" if status == "preparing" else "按计划跟进"
        action_sheet.write_row(
            row_index,
            0,
            [
                "高",
                action,
                _company_name(job),
                job.get("title"),
                job.get("deadline_at"),
                app.get("next_follow_up_at"),
                None,
            ],
            text,
        )
        _write_http_link(
            action_sheet,
            row_index,
            6,
            job.get("apply_url") or job.get("detail_url"),
            "打开",
            link,
        )
    _finish_table(
        action_sheet, len(actions), len(action_headers), [10, 20, 18, 22, 14, 20, 10]
    )

    dictionary_headers = ["字段", "含义", "填写规则"]
    dictionary = _sheet_frame(
        workbook, "数据字典", dictionary_headers, title, subtitle, header
    )
    dictionary_rows = [
        [
            "投递状态",
            "岗位当前求职阶段",
            "仅使用 considering / preparing / applied / assessment / interview / offer / rejected / withdrawn / expired",
        ],
        ["匹配分", "证据规则计算的 0-100 分", "仅供排序，不替代人工判断"],
        ["下次跟进", "计划提醒时间", "使用带时区的日期时间"],
        [
            "链接",
            "JD、投递入口、官网、报告和简历",
            "点击即可跳转；本地文件移动后需重新导出",
        ],
        ["事实边界", "简历所用个人信息", "只允许已确认且可用于简历的事实"],
    ]
    for row_index, row in enumerate(dictionary_rows, start=3):
        dictionary.write_row(row_index, 0, row, note)
    _finish_table(
        dictionary, len(dictionary_rows), len(dictionary_headers), [18, 28, 72]
    )
    workbook.close()


def _sheet_frame(
    workbook: Any,
    name: str,
    headers: Sequence[str],
    title_format: Any,
    subtitle_format: Any,
    header_format: Any,
) -> Any:
    sheet = workbook.add_worksheet(name)
    sheet.hide_gridlines(2)
    sheet.freeze_panes(3, 2)
    sheet.set_tab_color("#2F6B57")
    sheet.merge_range(0, 0, 0, len(headers) - 1, f"BioJob｜{name}", title_format)
    sheet.merge_range(
        1,
        0,
        1,
        len(headers) - 1,
        "本地生成 · 链接可点击 · 内容以应用数据库为准",
        subtitle_format,
    )
    sheet.write_row(2, 0, list(headers), header_format)
    sheet.set_row(0, 28)
    sheet.set_row(2, 24)
    return sheet


def _finish_table(
    sheet: Any, data_count: int, column_count: int, widths: Sequence[int]
) -> None:
    last_row = max(3, data_count + 2)
    sheet.autofilter(2, 0, last_row, column_count - 1)
    for index, width in enumerate(widths):
        sheet.set_column(index, index, width)
    sheet.set_default_row(19)


def _write_http_link(
    sheet: Any, row: int, column: int, url: Any, label: str, fmt: Any
) -> None:
    if isinstance(url, str) and url:
        sheet.write_url(row, column, url, fmt, label)


def _write_file_link(
    sheet: Any, row: int, column: int, path: Any, label: str, fmt: Any
) -> None:
    if isinstance(path, str) and path:
        normalized = Path(path).resolve().as_posix()
        sheet.write_url(row, column, f"external:{normalized}", fmt, label)


def _add_section_heading(document: DocumentObject, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(47, 107, 87)
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "86A899")
    border.append(bottom)
    paragraph._p.get_or_add_pPr().append(border)


def _display_value(value: Any) -> str:
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, list):
        return "；".join(_display_value(item) for item in value if _display_value(item))
    if isinstance(value, dict):
        return "；".join(
            f"{key}：{_display_value(item)}"
            for key, item in value.items()
            if _display_value(item)
        )
    if value is None:
        return ""
    return str(value)


def _first_fact_value(facts: Iterable[Mapping[str, Any]], keys: set[str]) -> str | None:
    for fact in facts:
        if str(fact.get("fact_key")) in keys:
            value = _display_value(fact.get("value"))
            if value:
                return value
    return None


def _company_name(job: Mapping[str, Any]) -> str:
    company = job.get("company")
    if isinstance(company, Mapping):
        return str(company.get("name") or company.get("canonical_name") or "")
    return str(job.get("company_name") or "")


def _safe_filename(value: str, *, fallback: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "-", value).strip(" .-")
    return (cleaned or fallback)[:80]
